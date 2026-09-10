import io
from collections.abc import AsyncIterator
from pathlib import Path
from uuid import UUID

import pytest
import xlwt  # type: ignore[import-untyped]
from openpyxl import Workbook  # type: ignore[import-untyped]

from enterprise_rag.adapters.loaders import SpreadsheetLoader
from enterprise_rag.domain import AppError, ErrorCode, RootKind
from enterprise_rag.ports import IngestionContext


class MemorySource:
    def __init__(self, data: bytes, *, name: str, media_type: str) -> None:
        self.data = data
        self.name = name
        self.media_type = media_type

    async def chunks(self) -> AsyncIterator[bytes]:
        midpoint = len(self.data) // 2
        yield self.data[:midpoint]
        yield self.data[midpoint:]


@pytest.fixture
def context(tmp_path: Path) -> IngestionContext:
    return IngestionContext(
        tenant_id=UUID("01900000-0000-7000-8000-000000000801"),
        document_id=UUID("01900000-0000-7000-8000-000000000802"),
        version_id=UUID("01900000-0000-7000-8000-000000000803"),
        temporary_directory=tmp_path,
    )


def make_xlsx() -> bytes:
    workbook = Workbook()
    first = workbook.active
    first.title = "产品"
    first.append([])
    first.append([None, "名称", "数量", "合计", None])
    first.append([None, "知识库", 2, "=B3*C3", None])
    first.append([])
    first.append([None, "检索器", 3, "=B5*C5", None])
    second = workbook.create_sheet("说明")
    second.append(["语言", "值"])
    second.append(["中文", "支持"])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def make_xls() -> bytes:
    workbook = xlwt.Workbook()
    worksheet = workbook.add_sheet("Legacy")
    worksheet.write(1, 1, "名称")
    worksheet.write(1, 2, "值")
    worksheet.write(2, 1, "旧格式")
    worksheet.write(2, 2, 7)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


@pytest.mark.anyio
async def test_xlsx_preserves_sheets_formulas_row_locations_and_repeated_headers(
    context: IngestionContext,
) -> None:
    roots = await SpreadsheetLoader(rows_per_root=1).load(
        MemorySource(
            make_xlsx(),
            name="产品.xlsx",
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ),
        context,
    )

    assert len(roots) == 3
    assert all(root.kind is RootKind.SHEET_ROWS for root in roots)
    assert [root.source_locator["sheet"] for root in roots] == ["产品", "产品", "说明"]
    assert roots[0].source_locator == {"sheet": "产品", "start_row": 3, "end_row": 3}
    assert roots[1].source_locator == {"sheet": "产品", "start_row": 5, "end_row": 5}
    assert roots[0].raw_text.startswith("| 名称 | 数量 | 合计 |")
    assert roots[1].raw_text.startswith("| 名称 | 数量 | 合计 |")
    assert "=B3*C3" in roots[0].raw_text
    assert roots[0].metadata["formula_cells"] == ({"row": 3, "column": 3, "formula": "=B3*C3"},)
    assert "| 中文 | 支持 |" in roots[2].raw_text


@pytest.mark.anyio
async def test_legacy_xls_is_parsed_independently(context: IngestionContext) -> None:
    roots = await SpreadsheetLoader().load(
        MemorySource(make_xls(), name="legacy.xls", media_type="application/vnd.ms-excel"),
        context,
    )

    assert len(roots) == 1
    assert roots[0].source_locator == {"sheet": "Legacy", "start_row": 3, "end_row": 3}
    assert "| 名称 | 值 |" in roots[0].raw_text
    assert "| 旧格式 | 7 |" in roots[0].raw_text


@pytest.mark.anyio
async def test_csv_supports_bom_chinese_long_tables_and_explicit_encoding_fallback(
    context: IngestionContext,
) -> None:
    csv_data = "名称,值\n" + "\n".join(f"项目{i},{i}" for i in range(61))
    roots = await SpreadsheetLoader(rows_per_root=60).load(
        MemorySource(b"\xef\xbb\xbf" + csv_data.encode(), name="data.csv", media_type="text/csv"),
        context,
    )
    assert len(roots) == 2
    assert all(root.raw_text.startswith("| 名称 | 值 |") for root in roots)
    assert roots[1].source_locator["start_row"] == 62

    gbk = "名称,值\n中文,可用".encode("gb18030")
    fallback = await SpreadsheetLoader(csv_fallback_encoding="gb18030").load(
        MemorySource(gbk, name="legacy.csv", media_type="text/csv"), context
    )
    assert "| 中文 | 可用 |" in fallback[0].raw_text
    assert fallback[0].source_locator["sheet"] == "CSV (gb18030)"


@pytest.mark.anyio
async def test_bad_encoding_empty_corrupt_type_and_close_are_stable(
    context: IngestionContext,
) -> None:
    loader = SpreadsheetLoader()
    cases = (
        (
            MemorySource(b"\xff", name="bad.csv", media_type="text/csv"),
            ErrorCode.DOCUMENT_ENCODING_INVALID,
        ),
        (MemorySource(b"\n,\n", name="empty.csv", media_type="text/csv"), ErrorCode.DOCUMENT_EMPTY),
        (
            MemorySource(b"bad", name="bad.xlsx", media_type="application/zip"),
            ErrorCode.DOCUMENT_CORRUPT,
        ),
        (
            MemorySource(b"a,b", name="bad.xls", media_type="text/csv"),
            ErrorCode.DOCUMENT_UNSUPPORTED_TYPE,
        ),
    )
    for source, code in cases:
        with pytest.raises(AppError) as raised:
            await loader.load(source, context)
        assert raised.value.code is code

    await loader.aclose()
    await loader.aclose()
    with pytest.raises(RuntimeError, match="closed"):
        await loader.load(MemorySource(b"a,b", name="x.csv", media_type="text/csv"), context)
