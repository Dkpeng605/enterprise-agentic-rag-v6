"""Deterministic XLSX, XLS, and CSV loader with repeated table headers."""

import csv
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date, datetime, time
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any

import xlrd  # type: ignore[import-untyped]
from openpyxl import load_workbook  # type: ignore[import-untyped]

from enterprise_rag.domain.documents import RootKind
from enterprise_rag.domain.errors import AppError, ErrorCode
from enterprise_rag.ports.loader import BinarySource, IngestionContext, LoadedRoot
from enterprise_rag.ports.provider import ProviderHealth, ProviderInfo, ProviderKind

_XLSX_MEDIA = frozenset(
    {
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/zip",
    }
)
_XLS_MEDIA = frozenset({"application/vnd.ms-excel", "application/xls"})
_CSV_MEDIA = frozenset({"text/csv", "application/csv", "text/plain"})
_XLS_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


class SpreadsheetLoaderError(AppError):
    """A stable client-safe spreadsheet parsing failure."""


@dataclass(frozen=True, slots=True)
class _Cell:
    value: str
    formula: str | None = None


@dataclass(frozen=True, slots=True)
class _Sheet:
    name: str
    rows: tuple[tuple[_Cell, ...], ...]
    row_numbers: tuple[int, ...]


class SpreadsheetLoader:
    def __init__(
        self, *, rows_per_root: int = 60, csv_fallback_encoding: str | None = None
    ) -> None:
        if not 1 <= rows_per_root <= 1_000:
            raise ValueError("rows_per_root must be between 1 and 1000")
        if csv_fallback_encoding is not None and not csv_fallback_encoding.strip():
            raise ValueError("csv_fallback_encoding must not be blank")
        self._rows_per_root = rows_per_root
        self._csv_fallback_encoding = csv_fallback_encoding
        self._closed = False

    def info(self) -> ProviderInfo:
        return ProviderInfo(
            kind=ProviderKind.LOADER,
            name="spreadsheets",
            version="1",
            capabilities=frozenset({"xlsx", "xls", "csv", "sheets", "formulas"}),
            is_remote=False,
            health=ProviderHealth.UNAVAILABLE if self._closed else ProviderHealth.HEALTHY,
        )

    def supports(self, media_type: str, suffix: str) -> bool:
        media = media_type.casefold().split(";", 1)[0].strip()
        normalized_suffix = suffix.casefold()
        return (
            (normalized_suffix == ".xlsx" and media in _XLSX_MEDIA)
            or (normalized_suffix == ".xls" and media in _XLS_MEDIA)
            or (normalized_suffix == ".csv" and media in _CSV_MEDIA)
        )

    async def load(self, source: BinarySource, context: IngestionContext) -> list[LoadedRoot]:
        del context
        if self._closed:
            raise RuntimeError("spreadsheet Loader is closed")
        suffix = Path(source.name).suffix.casefold()
        if not self.supports(source.media_type, suffix):
            raise SpreadsheetLoaderError(
                ErrorCode.DOCUMENT_UNSUPPORTED_TYPE,
                "The source is not a supported spreadsheet.",
                {"media_type": source.media_type, "suffix": suffix},
            )
        data = bytearray()
        async for chunk in source.chunks():
            if not isinstance(chunk, bytes):
                raise TypeError("source chunks must be bytes")
            data.extend(chunk)
        try:
            if suffix == ".xlsx":
                if not bytes(data).startswith(b"PK"):
                    raise ValueError("invalid XLSX signature")
                sheets = self._read_xlsx(bytes(data))
            elif suffix == ".xls":
                if not bytes(data).startswith(_XLS_SIGNATURE):
                    raise ValueError("invalid XLS signature")
                sheets = self._read_xls(bytes(data))
            else:
                sheets = (self._read_csv(bytes(data)),)
            roots = self._to_roots(sheets, suffix[1:])
        except AppError:
            raise
        except (csv.Error, OSError, ValueError, xlrd.XLRDError) as error:
            raise SpreadsheetLoaderError(
                ErrorCode.DOCUMENT_CORRUPT,
                "The spreadsheet could not be decoded.",
            ) from error
        if not roots:
            raise SpreadsheetLoaderError(
                ErrorCode.DOCUMENT_EMPTY,
                "The spreadsheet contains no non-empty cells.",
            )
        return roots

    async def aclose(self) -> None:
        self._closed = True

    @staticmethod
    def _read_xlsx(data: bytes) -> tuple[_Sheet, ...]:
        formulas = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        values = load_workbook(BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            sheets: list[_Sheet] = []
            for formula_sheet, value_sheet in zip(
                formulas.worksheets, values.worksheets, strict=True
            ):
                rows: list[tuple[_Cell, ...]] = []
                row_numbers: list[int] = []
                formula_rows = formula_sheet.iter_rows()
                value_rows = value_sheet.iter_rows()
                for row_number, (formula_row, value_row) in enumerate(
                    zip(formula_rows, value_rows, strict=True), start=1
                ):
                    cells: list[_Cell] = []
                    for formula_cell, value_cell in zip(formula_row, value_row, strict=True):
                        formula = (
                            str(formula_cell.value)
                            if formula_cell.data_type == "f" and formula_cell.value is not None
                            else None
                        )
                        raw_value = value_cell.value
                        if raw_value is None and formula is not None:
                            raw_value = formula
                        cells.append(_Cell(SpreadsheetLoader._format_value(raw_value), formula))
                    rows.append(tuple(cells))
                    row_numbers.append(row_number)
                sheets.append(SpreadsheetLoader._trim_sheet(formula_sheet.title, rows, row_numbers))
            return tuple(sheets)
        finally:
            formulas.close()
            values.close()

    @staticmethod
    def _read_xls(data: bytes) -> tuple[_Sheet, ...]:
        workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
        try:
            sheets: list[_Sheet] = []
            for sheet_name in workbook.sheet_names():
                worksheet = workbook.sheet_by_name(sheet_name)
                rows = [
                    tuple(
                        _Cell(SpreadsheetLoader._format_xls_cell(cell, workbook.datemode))
                        for cell in worksheet.row(row_index)
                    )
                    for row_index in range(worksheet.nrows)
                ]
                row_numbers = list(range(1, worksheet.nrows + 1))
                sheets.append(SpreadsheetLoader._trim_sheet(sheet_name, rows, row_numbers))
            return tuple(sheets)
        finally:
            workbook.release_resources()

    def _read_csv(self, data: bytes) -> _Sheet:
        try:
            text = data.decode("utf-8-sig")
            encoding = "utf-8"
        except UnicodeDecodeError as utf8_error:
            if self._csv_fallback_encoding is None:
                raise SpreadsheetLoaderError(
                    ErrorCode.DOCUMENT_ENCODING_INVALID,
                    "The CSV is not valid UTF-8 and no fallback encoding is configured.",
                    {"encoding": "utf-8"},
                ) from utf8_error
            try:
                text = data.decode(self._csv_fallback_encoding)
                encoding = self._csv_fallback_encoding
            except (LookupError, UnicodeDecodeError) as fallback_error:
                raise SpreadsheetLoaderError(
                    ErrorCode.DOCUMENT_ENCODING_INVALID,
                    "The CSV does not match its configured fallback encoding.",
                    {"encoding": self._csv_fallback_encoding},
                ) from fallback_error
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        parsed = list(csv.reader(StringIO(text, newline=""), dialect))
        rows = [tuple(_Cell(value) for value in row) for row in parsed]
        sheet = self._trim_sheet("CSV", rows, list(range(1, len(rows) + 1)))
        return _Sheet(f"{sheet.name} ({encoding})", sheet.rows, sheet.row_numbers)

    def _to_roots(self, sheets: Sequence[_Sheet], format_name: str) -> list[LoadedRoot]:
        roots: list[LoadedRoot] = []
        for sheet in sheets:
            if not sheet.rows:
                continue
            header = tuple(
                cell.value or f"Column {column + 1}" for column, cell in enumerate(sheet.rows[0])
            )
            data_rows = sheet.rows[1:]
            data_numbers = sheet.row_numbers[1:]
            chunks = [
                (
                    data_rows[index : index + self._rows_per_root],
                    data_numbers[index : index + self._rows_per_root],
                )
                for index in range(0, len(data_rows), self._rows_per_root)
            ] or [((), ())]
            for rows, row_numbers in chunks:
                formula_cells: list[dict[str, object]] = []
                for row_number, row in zip(row_numbers, rows, strict=True):
                    for column, cell in enumerate(row, start=1):
                        if cell.formula is not None:
                            formula_cells.append(
                                {"row": row_number, "column": column, "formula": cell.formula}
                            )
                start_row = row_numbers[0] if row_numbers else sheet.row_numbers[0]
                end_row = row_numbers[-1] if row_numbers else sheet.row_numbers[0]
                roots.append(
                    LoadedRoot(
                        ordinal=len(roots),
                        kind=RootKind.SHEET_ROWS,
                        source_locator={
                            "sheet": sheet.name,
                            "start_row": start_row,
                            "end_row": end_row,
                        },
                        raw_text=self._markdown_table(header, rows),
                        metadata={
                            "format": format_name,
                            "sheet": sheet.name,
                            "header_row": sheet.row_numbers[0],
                            "formula_cells": formula_cells,
                        },
                    )
                )
        return roots

    @staticmethod
    def _trim_sheet(
        name: str, rows: Sequence[tuple[_Cell, ...]], row_numbers: Sequence[int]
    ) -> _Sheet:
        non_empty = [
            (number, row)
            for number, row in zip(row_numbers, rows, strict=True)
            if any(cell.value.strip() for cell in row)
        ]
        if not non_empty:
            return _Sheet(name, (), ())
        occupied = [
            column for _, row in non_empty for column, cell in enumerate(row) if cell.value.strip()
        ]
        first_column, last_column = min(occupied), max(occupied)
        width = last_column - first_column + 1
        cropped = []
        for _, row in non_empty:
            padded = row + tuple(_Cell("") for _ in range(last_column + 1 - len(row)))
            cropped.append(padded[first_column : first_column + width])
        return _Sheet(
            name,
            tuple(cropped),
            tuple(number for number, _ in non_empty),
        )

    @staticmethod
    def _markdown_table(header: Sequence[str], rows: Sequence[Sequence[_Cell]]) -> str:
        def escaped(value: str) -> str:
            return value.replace("\n", "<br>").replace("|", "\\|")

        lines = [
            "| " + " | ".join(escaped(value) for value in header) + " |",
            "| " + " | ".join("---" for _ in header) + " |",
        ]
        lines.extend("| " + " | ".join(escaped(cell.value) for cell in row) + " |" for row in rows)
        return "\n".join(lines)

    @staticmethod
    def _format_value(value: object) -> str:
        if value is None:
            return ""
        if isinstance(value, (datetime, date, time)):
            return value.isoformat()
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value)

    @staticmethod
    def _format_xls_cell(cell: Any, datemode: int) -> str:
        if cell.ctype == xlrd.XL_CELL_DATE:
            converted: datetime = xlrd.xldate_as_datetime(cell.value, datemode)
            return converted.isoformat()
        if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
            return ""
        if cell.ctype == xlrd.XL_CELL_BOOLEAN:
            return "true" if cell.value else "false"
        return SpreadsheetLoader._format_value(cell.value)
